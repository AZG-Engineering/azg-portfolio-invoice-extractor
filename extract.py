"""Read every PDF invoice in a folder and write one spreadsheet row per invoice.

Usage:
    python extract.py <pdf_folder> <output.xlsx> [--vendors vendors.toml]

Each row holds Vendor, Invoice Date, Invoice Number and Total, plus the Source File
and a Review column. The extractor never guesses: a field it can't read with
certainty is left blank, and the Review column says why.
"""

from __future__ import annotations

import argparse
import re
import sys
import tomllib
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

AUTHOR = "AZG Engineering"
HEADERS = ("Vendor", "Invoice Date", "Invoice Number", "Total", "Source File", "Review")
COLUMN_WIDTHS = (26, 14, 18, 13, 24, 46)
# Colours from the AZG build style guide.
HEADER_FILL = "1F3A5F"  # navy header band, white text
REVIEW_FILL = "FDE7B0"  # amber: this row needs a person
REVIEW_TEXT = "7A4A00"
TEXT = "1F2937"
RULE = "E2E8F0"  # the one thin line under each row

NO_TEXT = "No text layer (scanned image): enter by hand"
VENDOR_KEYS = ("name", "match", "invoice_number", "invoice_date", "date_format", "total")


@dataclass
class Row:
    source_file: str
    vendor: str | None = None
    invoice_date: date | None = None
    invoice_number: str | None = None
    total: Decimal | None = None
    reasons: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.reasons

    @property
    def review(self) -> str:
        return "; ".join(self.reasons)


def load_vendors(path: Path) -> list[dict]:
    with open(path, "rb") as handle:
        vendors = tomllib.load(handle).get("vendor", [])
    if not vendors:
        raise ValueError(f"{path}: no [[vendor]] blocks found")
    for vendor in vendors:
        missing = [key for key in VENDOR_KEYS if key not in vendor]
        if missing:
            raise ValueError(f"{path}: vendor {vendor.get('name', '?')!r} is missing {', '.join(missing)}")
        for key in ("invoice_number", "invoice_date", "total"):
            if re.compile(vendor[key]).groups != 1:
                raise ValueError(f"{path}: {vendor['name']!r} {key} needs exactly one pair of brackets")
    return vendors


def read_text(pdf_path: Path) -> str:
    with pdfplumber.open(pdf_path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def single_match(pattern: str, text: str) -> tuple[str | None, str | None]:
    """Return (value, None) when exactly one distinct value matches, else (None, why)."""
    found = sorted({match.strip() for match in re.findall(pattern, text)})
    if not found:
        return None, "not found"
    if len(found) > 1:
        return None, "has more than one candidate (" + " / ".join(found) + ")"
    return found[0], None


def parse_invoice(text: str, vendors: list[dict], source_file: str) -> Row:
    row = Row(source_file)
    if not text.strip():
        row.reasons.append(NO_TEXT)
        return row

    lowered = text.lower()
    matching = [v for v in vendors if any(marker.lower() in lowered for marker in v["match"])]
    if not matching:
        row.reasons.append("Vendor not recognised: add its layout to vendors.toml")
        return row
    if len(matching) > 1:
        row.reasons.append("More than one vendor layout matches: " + ", ".join(v["name"] for v in matching))
        return row
    vendor = matching[0]
    row.vendor = vendor["name"]

    number, why = single_match(vendor["invoice_number"], text)
    if number is None:
        row.reasons.append(f"Invoice number {why}")
    else:
        row.invoice_number = number

    raw_date, why = single_match(vendor["invoice_date"], text)
    if raw_date is None:
        row.reasons.append(f"Invoice date {why}")
    else:
        try:
            row.invoice_date = datetime.strptime(raw_date, vendor["date_format"]).date()
        except ValueError:
            row.reasons.append(f"Invoice date unreadable ({raw_date})")

    raw_total, why = single_match(vendor["total"], text)
    if raw_total is None:
        row.reasons.append(f"Total {why}")
    else:
        try:
            row.total = Decimal(raw_total.replace(",", ""))
        except InvalidOperation:
            row.reasons.append(f"Total unreadable ({raw_total})")

    return row


def extract_folder(folder: Path, vendors: list[dict]) -> list[Row]:
    rows = []
    for pdf_path in sorted(folder.glob("*.pdf"), key=lambda p: p.name.lower()):
        try:
            text = read_text(pdf_path)
        except Exception as error:  # one unreadable file must not stop the batch
            rows.append(Row(pdf_path.name, reasons=[f"Could not open PDF ({type(error).__name__})"]))
            continue
        rows.append(parse_invoice(text, vendors, pdf_path.name))
    return rows


def summary(rows: list[Row]) -> str:
    ok = sum(1 for row in rows if row.ok)
    return f"{len(rows)} files, {ok} OK, {len(rows) - ok} need review"


def write_xlsx(rows: list[Row], out_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoices"
    workbook.properties.creator = AUTHOR
    workbook.properties.lastModifiedBy = AUTHOR
    workbook.properties.title = "Extracted invoices"

    # No boxes: one thin rule under each row, and colour only where it means something.
    border = Border(bottom=Side(style="thin", color=RULE))
    review_fill = PatternFill("solid", start_color=REVIEW_FILL)

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    for row in rows:
        sheet.append((
            row.vendor,
            row.invoice_date,
            row.invoice_number,
            float(row.total) if row.total is not None else None,
            row.source_file,
            row.review or None,
        ))
        cells = sheet[sheet.max_row]
        cells[1].number_format = "yyyy-mm-dd"
        cells[3].number_format = "#,##0.00"
        cells[5].alignment = Alignment(wrap_text=True, vertical="center")
        for cell in cells:
            cell.border = border
            cell.font = Font(name="Calibri", size=11, color=TEXT if row.ok else REVIEW_TEXT)
            if cell.alignment.vertical is None:
                cell.alignment = Alignment(vertical="center")
            if not row.ok:
                cell.fill = review_fill
        sheet.row_dimensions[sheet.max_row].height = 26

    for letter, width in zip("ABCDEF", COLUMN_WIDTHS):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pull vendor, date, number and total from a folder of PDF invoices.")
    parser.add_argument("pdf_folder", type=Path, help="folder holding the PDF invoices")
    parser.add_argument("output", type=Path, help="spreadsheet to write, e.g. output\\invoices.xlsx")
    parser.add_argument("--vendors", type=Path, default=Path(__file__).parent / "vendors.toml",
                        help="vendor layouts file (default: vendors.toml next to this script)")
    args = parser.parse_args(argv)

    if not args.pdf_folder.is_dir():
        print(f"Folder not found: {args.pdf_folder}", file=sys.stderr)
        return 2
    try:
        vendors = load_vendors(args.vendors)
    except (OSError, ValueError, tomllib.TOMLDecodeError) as error:
        print(f"Can't use the vendor layouts file: {error}", file=sys.stderr)
        return 2

    rows = extract_folder(args.pdf_folder, vendors)
    if not rows:
        print(f"No PDF files in {args.pdf_folder}", file=sys.stderr)
        return 1
    try:
        write_xlsx(rows, args.output)
    except PermissionError:
        print(f"Can't write {args.output}: close it in Excel and run again.", file=sys.stderr)
        return 1

    print(summary(rows))
    for row in rows:
        if not row.ok:
            print(f"  REVIEW  {row.source_file}: {row.review}")
    print(f"Saved {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
