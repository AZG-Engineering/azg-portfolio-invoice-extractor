"""Tests for the invoice extractor, run against the made-up invoices in samples/."""

import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import extract  # noqa: E402
import generate_samples  # noqa: E402

SAMPLES = ROOT / "samples"

# What a person reading each clean sample invoice would type in.
EXPECTED = {
    "contoso-0318.pdf": ("Contoso Freight", date(2026, 8, 11), "CF-2026-0318", Decimal("818.23")),
    "contoso-0331.pdf": ("Contoso Freight", date(2026, 8, 27), "CF-2026-0331", Decimal("2485.30")),
    "contoso-0346.pdf": ("Contoso Freight", date(2026, 9, 15), "CF-2026-0346", Decimal("383.52")),
    "fabrikam-00771.pdf": ("Fabrikam Office Goods", date(2026, 8, 6), "FOG-00771", Decimal("278.94")),
    "fabrikam-00784.pdf": ("Fabrikam Office Goods", date(2026, 9, 2), "FOG-00784", Decimal("405.07")),
    "northwind-10482.pdf": ("Northwind Supply Co.", date(2026, 8, 3), "NW-10482", Decimal("255.57")),
    "northwind-10497.pdf": ("Northwind Supply Co.", date(2026, 8, 19), "NW-10497", Decimal("258.96")),
    "northwind-10513.pdf": ("Northwind Supply Co.", date(2026, 9, 8), "NW-10513", Decimal("1071.13")),
}
MISSING_TOTAL = "northwind-10528.pdf"
SCANNED = "scan-0007.pdf"


@pytest.fixture(scope="module")
def vendors():
    return extract.load_vendors(ROOT / "vendors.toml")


@pytest.fixture(scope="module")
def rows(vendors):
    return {row.source_file: row for row in extract.extract_folder(SAMPLES, vendors)}


def test_every_sample_file_is_read(rows):
    assert sorted(rows) == sorted([*EXPECTED, MISSING_TOTAL, SCANNED])


@pytest.mark.parametrize("file_name", sorted(EXPECTED))
def test_clean_invoice_fields(rows, file_name):
    row = rows[file_name]
    assert (row.vendor, row.invoice_date, row.invoice_number, row.total) == EXPECTED[file_name]
    assert row.ok
    assert row.review == ""


def test_missing_total_is_left_blank_and_flagged(rows):
    row = rows[MISSING_TOTAL]
    assert row.total is None
    assert row.review == "Total not found"
    # The fields that can be read are still filled in.
    assert row.vendor == "Northwind Supply Co."
    assert row.invoice_number == "NW-10528"
    assert row.invoice_date == date(2026, 9, 22)


def test_scanned_invoice_is_all_blank_and_flagged(rows):
    row = rows[SCANNED]
    assert (row.vendor, row.invoice_date, row.invoice_number, row.total) == (None, None, None, None)
    assert "No text layer" in row.review


def test_summary_counts(rows):
    assert extract.summary(list(rows.values())) == "10 files, 8 OK, 2 need review"


def test_due_date_is_not_mistaken_for_invoice_date(rows):
    # Northwind invoices print a Due Date right under the Invoice Date.
    assert rows["northwind-10513.pdf"].invoice_date == date(2026, 9, 8)


# --- never guess -----------------------------------------------------------------

NORTHWIND_TEXT = (
    "Northwind Supply Co. INVOICE\n"
    "Bill To: Invoice No: NW-20001\n"
    "Example Client Co. Invoice Date: 09/30/2026\n"
    "Subtotal $100.00\n"
    "TOTAL DUE $106.00\n"
)


def test_text_fixture_parses_cleanly(vendors):
    row = extract.parse_invoice(NORTHWIND_TEXT, vendors, "x.pdf")
    assert row.ok
    assert row.total == Decimal("106.00")


def test_two_different_totals_are_not_guessed_between(vendors):
    row = extract.parse_invoice(NORTHWIND_TEXT + "TOTAL DUE $160.00\n", vendors, "x.pdf")
    assert row.total is None
    assert row.review == "Total has more than one candidate (106.00 / 160.00)"
    assert row.invoice_number == "NW-20001"


def test_same_total_printed_twice_is_accepted(vendors):
    row = extract.parse_invoice(NORTHWIND_TEXT + "TOTAL DUE $106.00\n", vendors, "x.pdf")
    assert row.ok
    assert row.total == Decimal("106.00")


def test_impossible_date_is_left_blank(vendors):
    row = extract.parse_invoice(NORTHWIND_TEXT.replace("09/30/2026", "13/45/2026"), vendors, "x.pdf")
    assert row.invoice_date is None
    assert row.review == "Invoice date unreadable (13/45/2026)"


def test_unknown_vendor_is_flagged_with_nothing_filled(vendors):
    row = extract.parse_invoice("Tailspin Toys\nInvoice 77\nTotal $12.00\n", vendors, "x.pdf")
    assert (row.vendor, row.invoice_date, row.invoice_number, row.total) == (None, None, None, None)
    assert row.review.startswith("Vendor not recognised")


def test_two_vendor_layouts_matching_is_flagged(vendors):
    row = extract.parse_invoice(NORTHWIND_TEXT + "CONTOSO FREIGHT\n", vendors, "x.pdf")
    assert row.vendor is None
    assert row.review.startswith("More than one vendor layout matches")


def test_broken_pdf_is_flagged_and_the_batch_continues(tmp_path, vendors):
    (tmp_path / "broken.pdf").write_bytes(b"this is not a pdf")
    (tmp_path / "good.pdf").write_bytes((SAMPLES / "contoso-0318.pdf").read_bytes())
    result = extract.extract_folder(tmp_path, vendors)
    assert [row.source_file for row in result] == ["broken.pdf", "good.pdf"]
    assert result[0].review.startswith("Could not open PDF")
    assert result[1].ok


# --- spreadsheet output ----------------------------------------------------------

def test_spreadsheet_contents_and_highlighting(tmp_path, rows):
    out = tmp_path / "out" / "invoices.xlsx"
    extract.write_xlsx([rows[name] for name in sorted(rows)], out)
    sheet = load_workbook(out)["Invoices"]

    assert tuple(cell.value for cell in sheet[1]) == extract.HEADERS
    assert sheet.max_row == 11

    by_file = {r[4].value: r for r in sheet.iter_rows(min_row=2)}
    clean = by_file["contoso-0331.pdf"]
    assert clean[0].value == "Contoso Freight"
    assert clean[1].value.date() == date(2026, 8, 27)
    assert clean[2].value == "CF-2026-0331"
    assert clean[3].value == pytest.approx(2485.30)
    assert clean[5].value is None
    assert clean[0].fill.fill_type is None

    for name in (MISSING_TOTAL, SCANNED):
        flagged = by_file[name]
        assert flagged[5].value
        assert all(cell.fill.start_color.rgb.endswith(extract.REVIEW_FILL) for cell in flagged)
    assert by_file[MISSING_TOTAL][3].value is None
    assert by_file[SCANNED][0].value is None


def test_spreadsheet_author_is_the_company(tmp_path, rows):
    out = tmp_path / "invoices.xlsx"
    extract.write_xlsx(list(rows.values()), out)
    properties = load_workbook(out).properties
    assert properties.creator == "AZG Engineering"
    assert properties.lastModifiedBy == "AZG Engineering"


# --- command line ----------------------------------------------------------------

def test_command_line_prints_the_summary(tmp_path, capsys):
    out = tmp_path / "invoices.xlsx"
    assert extract.main([str(SAMPLES), str(out)]) == 0
    printed = capsys.readouterr().out.splitlines()
    assert printed[0] == "10 files, 8 OK, 2 need review"
    assert any(MISSING_TOTAL in line and "Total not found" in line for line in printed)
    assert any(SCANNED in line for line in printed)
    assert out.exists()


def test_command_line_rejects_a_missing_folder(tmp_path, capsys):
    assert extract.main([str(tmp_path / "nope"), str(tmp_path / "x.xlsx")]) == 2
    assert "Folder not found" in capsys.readouterr().err


def test_command_line_reports_an_empty_folder(tmp_path, capsys):
    assert extract.main([str(tmp_path), str(tmp_path / "x.xlsx")]) == 1
    assert "No PDF files" in capsys.readouterr().err


# --- sample generator ------------------------------------------------------------

def test_samples_can_be_regenerated_identically(tmp_path, vendors, rows):
    generate_samples.generate(tmp_path)
    again = {row.source_file: row for row in extract.extract_folder(tmp_path, vendors)}
    assert again == rows


def test_sample_pdfs_carry_the_company_as_author():
    import pdfplumber

    for pdf_path in SAMPLES.glob("*.pdf"):
        with pdfplumber.open(pdf_path) as pdf:
            assert pdf.metadata.get("Author") == "AZG Engineering", pdf_path.name
